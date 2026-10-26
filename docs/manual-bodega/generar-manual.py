#!/usr/bin/env python3
"""Genera el manual Word de Bodega y Despacho a partir de capturas reales de SeePOS."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

DIR = Path(__file__).resolve().parent
OUT = DIR / "Manual-SeePOS-Bodega-y-Despacho.docx"

AZUL = RGBColor(0x10, 0x72, 0xA9)
NARANJA = RGBColor(0xEE, 0x75, 0x19)
GRIS = RGBColor(0x5C, 0x5C, 0x5C)
OSCURO = RGBColor(0x1F, 0x2A, 0x37)


def set_run(run, *, size=11, bold=False, color=OSCURO, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run(run, size=16 if level == 1 else 13, bold=True, color=AZUL if level == 1 else NARANJA)
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "1072A9")
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


def add_body(doc, text, *, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=11, bold=bold, italic=italic, color=OSCURO)
    return p


def add_bullet(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run(r1, size=11, bold=True)
        r2 = p.add_run(text)
        set_run(r2, size=11)
    else:
        r = p.add_run(text)
        set_run(r, size=11)
    return p


def add_callout(doc, titulo, texto, fill="EEF6FB"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.3)
    shade(p, fill)
    r = p.add_run(titulo)
    set_run(r, size=11, bold=True, color=AZUL)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.left_indent = Cm(0.3)
    shade(p2, fill)
    r2 = p2.add_run(texto)
    set_run(r2, size=10.5, color=OSCURO)
    return p2


def add_step(doc, numero, titulo, detalle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"Paso {numero}. ")
    set_run(r1, size=12, bold=True, color=NARANJA)
    r2 = p.add_run(titulo)
    set_run(r2, size=12, bold=True, color=OSCURO)
    add_body(doc, detalle)


def add_figure(doc, filename, caption):
    path = DIR / filename
    if not path.exists():
        add_body(doc, f"[Falta captura {filename}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(16.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    set_run(r, size=9.5, italic=True, color=GRIS)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rh = header.add_run("SeePOS  ·  Manual operativo")
    set_run(rh, size=9, color=GRIS)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = footer.add_run("Capturas tomadas del sistema real  ·  No incluye claves de acceso")
    set_run(rf, size=8.5, color=GRIS)

    banda = doc.add_paragraph()
    shade(banda, "1072A9")
    rb = banda.add_run(" SEEPOS")
    set_run(rb, size=14, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    t = doc.add_paragraph()
    t.paragraph_format.space_before = Pt(18)
    t.paragraph_format.space_after = Pt(4)
    rt = t.add_run("Bodega y despacho")
    set_run(rt, size=28, bold=True, color=AZUL)

    st = doc.add_paragraph()
    st.paragraph_format.space_after = Pt(6)
    rs = st.add_run("Cola de alistamiento, empaque, manifiesto, entrega y recibo de devoluciones")
    set_run(rs, size=14, color=NARANJA)

    add_body(
        doc,
        "Este manual muestra el flujo real en SeePOS, de la factura al camión. "
        "Las imágenes son pantallas del sistema. El alistamiento confirma una salida que ya se hizo al facturar: "
        "no vuelve a descontar inventario.",
    )

    add_callout(
        doc,
        "Ejemplo usado en las capturas",
        "Factura crédito n.º 6 (venta 52)  ·  10 unidades de Alimento Barf Perro Adulto (código 002).\n"
        "Bodega Facturación  ·  bulto BULT-1 de 1,25 kg  ·  "
        "ruta San José Centro, vehículo SJO-001, chofer Carlos Pérez.",
    )

    add_callout(
        doc,
        "Sobre las capturas",
        "Las imágenes se tomaron antes de las últimas mejoras (rutas con zonas, ingreso manual con supervisor, "
        "manifiesto editable, comprobante de entrega y recibo de mercadería). El texto de cada sección sí está actualizado; "
        "las pantallas nuevas se describen sin captura.",
        fill="FFF4EA",
    )

    add_heading_custom(doc, "1. Qué cubre el módulo Bodega", 1)
    add_body(
        doc,
        "El menú Bodega agrupa cinco pantallas. Se recorren en este orden operativo:",
    )
    add_bullet(doc, "rutas de despacho con las zonas que cubren, más vehículos y choferes. Solo el encargado de bodega crea rutas. Distinto de la ruta comercial del vendedor (comisiones).", bold_prefix="Rutas de entrega. ")
    add_bullet(doc, "las facturas, tiquetes, consignaciones y preventas pendientes de recoger. Entran todas las series, salvo la de Ferias.", bold_prefix="Cola de alistamiento. ")
    add_bullet(doc, "la mercadería devuelta por clientes: el inventario sube cuando bodega confirma lo que recibió.", bold_prefix="Recibo de mercadería. ")
    add_bullet(doc, "una orden a la vez. Un disparo de pistola = una caja o etiqueta. El lote es el de la factura.", bold_prefix="Empaque. ")
    add_bullet(doc, "armar el manifiesto con bultos empacados, escanear al camión y salir a ruta.", bold_prefix="Despacho. ")
    add_figure(doc, "01-menu-bodega.png", "Figura 1. Menú Bodega: Cola, Empaque, Despacho y Rutas de entrega (la captura es anterior: ya no incluye Ubicaciones, que se retiró).")

    add_heading_custom(doc, "2. Catálogos que se configuran una vez", 1)
    add_step(
        doc,
        1,
        "Registrar zona, vehículo y chofer",
        "Bodega → Rutas de entrega. Son tres listas independientes. La ruta de despacho no es la ruta comercial del SAC. "
        "En el ejemplo: ruta San José Centro, placa SJO-001, chofer Carlos Pérez.\n"
        "Cada ruta tiene un nombre, una casilla Activa y las zonas que cubre: una provincia completa, o un cantón o distrito concreto; "
        "puede ser de varias provincias o de una zona específica. Solo el encargado de bodega (permiso Crear/Editar) crea y modifica rutas; "
        "el despachador solo las elige en Despacho, y solo ve las activas. Una ruta inactiva ya no se ofrece pero no se pierde.",
    )
    add_figure(doc, "05-rutas.png", "Figura 2. Rutas de entrega: zonas, vehículos y choferes.")

    add_heading_custom(doc, "3. Cola de alistamiento", 1)
    add_body(
        doc,
        "Al emitir una factura (contado o crédito), un tiquete, una consignación o guardar una preventa, la orden entra sola en Pendiente. "
        "Entran todas las series, salvo la marcada «Sin alistamiento» en Parámetros → Series (la de Ferias, de entrega inmediata). "
        "También entra el envío de mercadería a consignación: al registrar una boleta de ingreso en Consignación, la orden aparece como «envío consignación» "
        "y se alista en la bodega central. Anular la boleta antes de que salga cancela la orden. "
        "Si falta mercadería, SeePOS anula la boleta y emite otra por lo empacado; si no se entrega, anula la boleta y lo empacado vuelve a la central "
        "(en merma en tránsito, además se da de baja). "
        "Si una factura ya emitida no apareció, use Traer factura con el número o el id de la venta y pulse Crear.\n"
        "La cola solo muestra las órdenes de la sucursal de su sesión. La columna Espera indica cuánto lleva la orden desde la facturación "
        "(o cuánto tardó bodega en iniciarla): verde, amarillo o rojo según los minutos configurados. "
        "Puede filtrar por estado, tipo, prioridad (alta = crédito) y documento (factura o tiquete).",
    )
    add_figure(
        doc,
        "02-cola-vacia.png",
        "Figura 3. Cola vacía (captura de la versión anterior).",
    )
    add_step(
        doc,
        2,
        "Traer una factura a la cola",
        "En Traer factura escriba el id de la venta (en el ejemplo, 52) o el número de documento y pulse Crear. "
        "SeePOS abre Empaque, o deja la orden en la cola para pulsar Empacar.",
    )
    add_figure(doc, "06-cola-traer-factura.png", "Figura 4. Traer factura: id 52 (factura crédito n.º 6).")
    add_figure(
        doc,
        "07-cola-con-orden.png",
        "Figura 5. La orden aparece con prioridad Alta (es crédito), estado Pendiente y avance 0 / 1 · 0,00 / 10,00.",
    )

    add_callout(
        doc,
        "Estados de la cola",
        "Pendiente → En alistamiento → Empacado → En ruta → Entregado (o No entregado). "
        "Anular en cola o en empaque no toca la factura: solo cancela el picking.",
        fill="FFF4EA",
    )

    add_heading_custom(doc, "4. Empaque", 1)
    add_step(
        doc,
        3,
        "Abrir la orden y aceptarla",
        "Desde la cola pulse Empacar. Verá el cliente/documento, las líneas, "
        "cantidad facturada y lote de la factura. Pulse Aceptar: la orden pasa a En alistamiento y se abre un bulto (BULT-n).",
    )
    add_figure(
        doc,
        "08-empaque-pendiente.png",
        "Figura 6. Orden pendiente. 10 unidades, lote de la factura 0. Botón Aceptar.",
    )
    add_figure(
        doc,
        "09-empaque-alistamiento.png",
        "Figura 7. En alistamiento. Campo Pistola / código, Cerrar bulto, Faltante y Marcar empacado. BULT-1 Abierto.",
    )

    add_step(
        doc,
        4,
        "Escanear (un disparo = una caja o etiqueta)",
        "Pistola USB o colectora en Pistola / código. El lote escaneado tiene que ser el de la factura. "
        "Una etiqueta GS1 puede traer lote (AI 10) y cantidad (AI 37) y cubre la línea de un disparo. "
        "En el ejemplo, un GS1 con cantidad 10 dejó la línea Completa (10,000 / 10,000). "
        "Se bloquea un lote vencido (el del código o el de la factura) o con menos vida útil de la que exige el cliente.\n"
        "Si la barra está dañada, o la etiqueta no trae lote, use Ingreso manual (supervisor): escriba el código del artículo y pida "
        "al supervisor su clave interna. Solo vale un usuario cuyo perfil tenga la opción «Supervisor de bodega». "
        "En el ingreso manual el lote es siempre el de la factura (no se puede elegir otro) y el escaneo queda registrado con quién autorizó. "
        "Los artículos sin ningún código de barras registrado se siguen identificando por su código, sin supervisor.",
    )
    add_step(
        doc,
        5,
        "Cerrar el bulto y marcar empacado",
        "Indique el peso en kg y pulse Cerrar bulto: el bulto queda Cerrado (BULT-1) y se abre la etiqueta para imprimir "
        "(A4 u 80 mm desde la lista de bultos). Al completar el 100 % de lo facturado, SeePOS cierra el bulto abierto y abre la etiqueta solo. "
        "Cuando todas las líneas están Completas, pulse Marcar empacado: los bultos que seguían abiertos se cierran y se imprimen. "
        "SeePOS avisa: pase a Despacho para armar el manifiesto.",
    )
    add_figure(
        doc,
        "10-empaque-bulto.png",
        "Figura 8. Línea completa, BULT-1 Cerrado 1,25 kg, escaneo GS1 con lote 0 y cantidad 10.",
    )
    add_figure(doc, "11-cola-empacado.png", "Figura 9. Cola filtrada en Empacado: 1 / 1 · 10,00 / 10,00.")
    add_figure(
        doc,
        "12-empaque-empacado.png",
        "Figura 10. Orden Empacado. Ya se puede confirmar Entregado / No entregado, o pasar a Despacho.",
    )

    add_callout(
        doc,
        "Faltante",
        "Si no se encuentra todo, pulse Faltante. En factura se emite una nota de crédito por lo no empacado. "
        "En preventa se quita la cantidad y se devuelve al inventario. Después puede marcar empacado el resto. "
        "No se usa si la línea ya está completa.",
        fill="FFF4EA",
    )

    add_heading_custom(doc, "5. Despacho", 1)
    add_step(
        doc,
        6,
        "Armar el manifiesto",
        "Bodega → Despacho. Marque los bultos empacados sin manifiesto (BULT-1 · factura 6). "
        "Elija la ruta (obligatoria; solo aparecen las activas), el vehículo y el chofer. Pulse Crear manifiesto. SeePOS abre el PDF del manifiesto.\n"
        "Cada bulto muestra la zona y la dirección de su cliente. Con una ruta elegida puede marcar «Solo bultos de la ruta» para ver únicamente "
        "los que caen en sus zonas; si carga un bulto fuera de las zonas, SeePOS lo avisa pero no lo impide: la decisión es del despachador. "
        "Si no hay rutas activas, pida al encargado de bodega que cree una.",
    )
    add_figure(
        doc,
        "13-despacho-armar.png",
        "Figura 11. BULT-1 seleccionado, ruta San José Centro, placa SJO-001, chofer Carlos Pérez.",
    )
    add_step(
        doc,
        7,
        "Escanear al camión y salir a ruta",
        "En el manifiesto pulse el campo Escanear bulto al camión y dispare el código BULT-n (o escríbalo). "
        "Hay que cargar todos los bultos del manifiesto, y el manifiesto debe tener ruta. Luego pulse Salir a ruta. La orden pasa a En ruta.",
    )
    add_callout(
        doc,
        "Corregir un manifiesto antes de salir",
        "Mientras el manifiesto esté Armado se puede cambiar la ruta, el vehículo y el chofer (Guardar cambios), quitar un bulto "
        "(queda libre para otro manifiesto) o anular el manifiesto con un motivo. Una vez que sale a ruta ya no se puede corregir.",
        fill="FFF4EA",
    )
    add_figure(doc, "14-cola-en-ruta.png", "Figura 12. Cola filtrada en En ruta. Documento 6 crédito, avance completo.")
    add_figure(
        doc,
        "15-empaque-en-ruta.png",
        "Figura 13. Empaque con la orden En ruta. BULT-1 En ruta. Botones Entregado y No entregado.",
    )

    add_heading_custom(doc, "6. Entrega", 1)
    add_step(
        doc,
        8,
        "Confirmar entregado",
        "Desde Empaque (o desde Despacho, factura por factura) pulse Entregado. "
        "Indique quién recibió. La firma en el recuadro es opcional: rastro de quién recibió, no certificado legal. "
        "Pulse Confirmar. Con la orden cerrada puede imprimir el Comprobante de entrega (qué se entregó, quién recibió, ruta, vehículo y chofer, "
        "y una línea para la firma a mano) desde Empaque o desde el manifiesto en Despacho.",
    )
    add_figure(
        doc,
        "16-confirmar-entrega.png",
        "Figura 14. Confirmar entrega: Recibido por María Rodríguez. Firma opcional.",
    )
    add_figure(
        doc,
        "17-cola-entregado.png",
        "Figura 15. Cola filtrada en Entregado. La orden quedó cerrada (1 / 1 · 10,00 / 10,00).",
    )

    add_callout(
        doc,
        "No entregado",
        "Si el producto vuelve o se pierde: pulse No entregado. "
        "Elija Volvió a bodega (reingreso) o Merma o pérdida en tránsito. "
        "En factura se emite nota de crédito por lo empacado; en preventa se suelta la reserva. "
        "El inventario no se vuelve a descontar (ya salió al facturar).",
        fill="FFF4EA",
    )

    add_heading_custom(doc, "7. Recibo de mercadería (devoluciones)", 1)
    add_body(
        doc,
        "Cuando se emite una devolución de venta con mercadería (nota de crédito o devolución interna), el inventario no sube de inmediato: "
        "la devolución entra a Bodega → Recibo de mercadería como una orden Pendiente de recibo.",
    )
    add_step(
        doc,
        9,
        "Recibir la mercadería",
        "Abra la orden con Recibir. Verá el cliente, la factura, las líneas con su lote y bodega, y la cantidad devuelta. "
        "Escriba la cantidad realmente recibida de cada línea (por defecto, lo devuelto; no puede ser mayor) y pulse Confirmar recibo. "
        "Solo entonces el inventario sube, y por lo recibido. Si se recibe menos de lo devuelto, la diferencia queda anotada como faltante.",
    )
    add_callout(
        doc,
        "Anular y excepciones",
        "Una orden pendiente se puede anular con un motivo (Anular orden): la mercadería no entra al inventario. "
        "Si se anula la devolución de origen mientras la orden sigue pendiente, la orden se anula sola; si ya se recibió, se revierte lo recibido. "
        "Los faltantes de alistamiento no pasan por recibo: esa mercadería nunca salió de bodega. "
        "Con Devoluciones:ExigirRecibo en false el inventario vuelve a subir al emitir la devolución.",
        fill="FFF4EA",
    )

    add_heading_custom(doc, "8. Resumen para el mostrador de bodega", 1)
    add_bullet(doc, "Las rutas de entrega (con sus zonas), los vehículos y los choferes se mantienen en catálogo; las rutas las crea el encargado de bodega.", bold_prefix="Antes. ")
    add_bullet(doc, "La venta entra sola, de cualquier serie salvo Ferias. Si no, Traer factura con el número o el id.", bold_prefix="Cola. ")
    add_bullet(doc, "Aceptar → escanear (lote de la factura; manual solo con supervisor) → la etiqueta sale al cerrar el bulto → marcar empacado.", bold_prefix="Empaque. ")
    add_bullet(doc, "Elegir bultos + ruta/vehículo/chofer → crear manifiesto (se puede corregir) → escanear BULT-n → salir a ruta.", bold_prefix="Despacho. ")
    add_bullet(doc, "Entregado pide quién recibió y permite imprimir el comprobante. No entregado corrige el documento (NC o preventa).", bold_prefix="Cierre. ")
    add_bullet(doc, "Las devoluciones se reciben en Recibo de mercadería; el inventario sube por lo recibido.", bold_prefix="Devoluciones. ")

    add_heading_custom(doc, "9. Qué no hacer", 1)
    add_bullet(doc, "Ajustar existencias otra vez al empacar: el stock ya salió con la factura.")
    add_bullet(doc, "Cambiar el lote en empaque. El lote es el de la factura, no se elige por FEFO.")
    add_bullet(doc, "Pedir a un compañero sin perfil de supervisor que «autorice» el ingreso manual: la clave debe ser de un supervisor de bodega.")
    add_bullet(doc, "Salir a ruta sin ruta o sin escanear todos los bultos del manifiesto.")
    add_bullet(doc, "Confundir la ruta de despacho (a dónde va el camión) con la ruta comercial de Facturación (la de las comisiones).")
    add_bullet(doc, "Confirmar Entregado sin el nombre de quien recibió.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(
        "Documento generado a partir de capturas reales de SeePOS. "
        "La factura crédito n.º 6 y la ruta son de demostración para este manual."
    )
    set_run(r, size=9.5, italic=True, color=GRIS)

    doc.save(OUT)
    print("WROTE", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    main()
