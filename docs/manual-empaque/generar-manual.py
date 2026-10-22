#!/usr/bin/env python3
"""Genera el manual Word de unidad vs caja a partir de capturas reales de SeePOS."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor

DIR = Path(__file__).resolve().parent
OUT = DIR / "Manual-SeePOS-Unidad-y-Caja.docx"

AZUL = RGBColor(0x10, 0x72, 0xA9)
NARANJA = RGBColor(0xEE, 0x75, 0x19)
GRIS = RGBColor(0x5C, 0x5C, 0x5C)
OSCURO = RGBColor(0x1F, 0x2A, 0x37)


def set_run(run, *, size=11, bold=False, color=OSCURO, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def shade(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)


def add_heading_custom(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run(run, size=16 if level == 1 else 13, bold=True, color=AZUL if level == 1 else NARANJA)
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "1072A9")
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


def add_body(doc, text, *, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=11, bold=bold, italic=italic, color=OSCURO)
    return p


def add_bullet(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run(r1, size=11, bold=True)
        r2 = p.add_run(text)
        set_run(r2, size=11)
    else:
        r = p.add_run(text)
        set_run(r, size=11)
    return p


def add_callout(doc, titulo, texto, fill="EEF6FB"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.3)
    shade(p, fill)
    r = p.add_run(titulo)
    set_run(r, size=11, bold=True, color=AZUL)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.left_indent = Cm(0.3)
    shade(p2, fill)
    r2 = p2.add_run(texto)
    set_run(r2, size=10.5, color=OSCURO)
    return p2


def add_step(doc, numero, titulo, detalle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"Paso {numero}. ")
    set_run(r1, size=12, bold=True, color=NARANJA)
    r2 = p.add_run(titulo)
    set_run(r2, size=12, bold=True, color=OSCURO)
    add_body(doc, detalle)


def add_figure(doc, filename, caption):
    path = DIR / filename
    if not path.exists():
        add_body(doc, f"[Falta captura {filename}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(16.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    set_run(r, size=9.5, italic=True, color=GRIS)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rh = header.add_run("SeePOS  ·  Manual operativo")
    set_run(rh, size=9, color=GRIS)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = footer.add_run("Capturas tomadas del sistema real  ·  No incluye claves de acceso")
    set_run(rf, size=8.5, color=GRIS)

    # Portada
    banda = doc.add_paragraph()
    banda.alignment = WD_ALIGN_PARAGRAPH.LEFT
    shade(banda, "1072A9")
    rb = banda.add_run(" SEEPOS")
    set_run(rb, size=14, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    t = doc.add_paragraph()
    t.paragraph_format.space_before = Pt(18)
    t.paragraph_format.space_after = Pt(4)
    rt = t.add_run("Unidad y caja")
    set_run(rt, size=28, bold=True, color=AZUL)

    st = doc.add_paragraph()
    st.paragraph_format.space_after = Pt(6)
    rs = st.add_run("Presentación empaquetada: cómo se configura y cómo se vende")
    set_run(rs, size=14, color=NARANJA)

    add_body(
        doc,
        "Este manual muestra el flujo real en SeePOS. Las imágenes son pantallas del sistema, "
        "no esquemas. El ejemplo usa un par de artículos de demostración: una tableta (unidad) "
        "y su caja de 10 tabletas.",
    )

    add_callout(
        doc,
        "Ejemplo usado en las capturas",
        "MAN-UND  —  Manual empaque — tableta (unidad)  ·  precio ₡113  ·  25 unidades en bodega Facturación.\n"
        "MAN-CAJ  —  Manual empaque — caja x 10 tabletas  ·  precio ₡1.017  ·  1 caja descuenta 10 unidades de MAN-UND.",
    )

    # 1
    add_heading_custom(doc, "1. Qué es una presentación empaquetada", 1)
    add_body(
        doc,
        "En SeePOS la unidad y la caja son dos artículos distintos (dos códigos). El inventario "
        "físico vive solo en la unidad. La caja no tiene existencias propias: al vender una caja "
        "el sistema rebaja el factor de unidades del artículo ligado.",
    )
    add_bullet(doc, "Dos códigos de artículo: uno para la unidad y otro para la caja o paquete.", bold_prefix=None)
    add_bullet(doc, "El stock se ajusta y se toma en el artículo unidad.", bold_prefix=None)
    add_bullet(doc, "La caja declara, en Otros datos, qué artículo rebaja y cuántas unidades trae cada caja.", bold_prefix=None)
    add_bullet(doc, "En el comprobante se imprime el código de la caja (con su precio). La descarga ocurre en la unidad.", bold_prefix=None)

    add_callout(
        doc,
        "Ojo: «Tipo de presentación» no activa el empaque",
        "El campo Tipo de presentación (UNIDAD, CAJA, etc.) es solo una etiqueta de catálogo. "
        "Lo que realmente convierte un código en caja es la pestaña Otros datos: Artículo a rebajar "
        "más Unidades por caja mayor que 1.",
        fill="FFF4EA",
    )

    # 2
    add_heading_custom(doc, "2. El par en Inventarios", 1)
    add_step(
        doc,
        1,
        "Abrir Inventarios",
        "Menú Inicio → Inventarios. En Código o descripción escriba parte del nombre (en el ejemplo: «Manual empaque») y pulse la lupa.",
    )
    add_figure(
        doc,
        "01-par-unidad-caja.png",
        "Figura 1. Los dos códigos aparecen juntos. La columna Existencia de la caja muestra 0 porque el stock físico está en la unidad (25).",
    )
    add_body(
        doc,
        "Esa columna 0 en la caja no significa que no se pueda vender. Significa que SeePOS no guarda "
        "un inventario paralelo de cajas. Las cajas disponibles se calculan: 25 unidades ÷ 10 = 2 cajas.",
    )

    # 3
    add_heading_custom(doc, "3. El artículo unidad", 1)
    add_step(
        doc,
        2,
        "Revisar datos generales de la unidad",
        "Pulse el lápiz de MAN-UND. Complete Código, Descripción, Impuesto, Familia y Código CABYS. "
        "Tipo de presentación puede quedar en UNIDAD. El aviso al pie indica que la existencia se maneja en la pestaña Existencias.",
    )
    add_figure(doc, "02-unidad-generales.png", "Figura 2. Artículo unidad — pestaña Datos generales.")

    add_step(
        doc,
        3,
        "Cargar existencias en la unidad",
        "Pestaña Existencias. Ahí sí se ajusta el inventario. En el ejemplo hay 25,00 unidades en la bodega Facturación. "
        "Para corregir, escriba el valor en «Ajustar a» y pulse Aplicar. Nunca haga este ajuste sobre el código de la caja.",
    )
    add_figure(doc, "03-unidad-existencias.png", "Figura 3. El stock vive en la unidad: 25 tabletas en Facturación.")

    # 4
    add_heading_custom(doc, "4. El artículo caja", 1)
    add_step(
        doc,
        4,
        "Crear el código de la caja",
        "En Inventarios pulse Agregar. Llene Datos generales como cualquier artículo: código propio (MAN-CAJ), "
        "descripción («caja x 10 tabletas»), familia, impuesto y precio de venta de la caja. "
        "Tipo de presentación puede seguir diciendo UNIDAD: no es lo que define el empaque.",
    )
    add_figure(doc, "04-caja-generales.png", "Figura 4. Artículo caja — Datos generales. Tiene código y precio propios.")

    add_step(
        doc,
        5,
        "Configurar el empaque en Otros datos",
        "Abra la pestaña Otros datos. En Artículo a rebajar elija el código unidad (MAN-UND). "
        "En Unidades por caja escriba el factor (10). Un valor mayor que 1 activa el comportamiento de caja. "
        "Guarde.",
    )
    add_figure(
        doc,
        "07-nuevo-otros.png",
        "Figura 5. Pestaña Otros datos al crear un artículo: Artículo a rebajar vacío y Unidades por caja en 0.",
    )
    add_figure(
        doc,
        "05-caja-otros.png",
        "Figura 6. MAN-CAJ ya configurado: Artículo a rebajar = MAN-UND, Unidades por caja = 10.",
    )

    add_step(
        doc,
        6,
        "Ver las cajas disponibles",
        "En la pestaña Existencias de la caja SeePOS muestra el cálculo, no un stock independiente. "
        "En el ejemplo: Existencia total 2,00 — «Cajas disponibles (unidades ÷ 10 de MAN-UND). El ajuste se hace en el artículo unidad». "
        "La columna Ajustar a está reemplazada por el enlace «Ajuste en el artículo unidad».",
    )
    add_figure(
        doc,
        "06-caja-existencias.png",
        "Figura 7. Existencias de la caja: 2,00 cajas calculadas (25 ÷ 10). No se ajusta desde aquí.",
    )

    add_callout(
        doc,
        "Reglas que el sistema aplica",
        "El factor debe ser mayor que 1. El artículo a rebajar no puede ser el mismo código. "
        "No se encadenan cajas (una caja no puede rebajar otra caja). "
        "No se hace toma física ni ajuste absoluto sobre el código empaquetado.",
        fill="FFF4EA",
    )

    # 5
    add_heading_custom(doc, "5. Facturar unidad y caja", 1)
    add_body(
        doc,
        "Para vender hace falta caja abierta y activar la venta con la clave interna. "
        "La clave interna es la misma contraseña con la que el usuario ingresó al sistema.",
    )

    add_step(
        doc,
        7,
        "Comprobar que hay caja abierta",
        "Menú Caja → Apertura Caja. Active la pantalla con la clave interna. "
        "Si el cajero ya tiene apertura, SeePOS avisa y no deja abrir otra hasta arquear y cerrar.",
    )
    add_figure(
        doc,
        "09-apertura-activada.png",
        "Figura 8. Apertura de caja activada. Este cajero ya tiene la apertura #7 abierta.",
    )

    add_step(
        doc,
        8,
        "Abrir Facturación y activar la venta",
        "Menú Inicio → Facturación (o la pestaña Venta). El banner pide la clave interna. "
        "Escríbala y pulse Activar. Use la misma contraseña del ingreso. "
        "Cuando es correcta aparece: «la venta está activado. Cajero: … · Apertura #n».",
    )
    add_figure(
        doc,
        "10-factura-bloqueada.png",
        "Figura 9. Facturación bloqueada: el formulario está inactivo hasta validar la clave interna.",
    )
    add_figure(
        doc,
        "11-factura-activada.png",
        "Figura 10. Venta activada. Ya se puede buscar cliente y artículos. Bodega: Facturación.",
    )

    add_step(
        doc,
        9,
        "Elegir el cliente",
        "En Buscar cliente escriba cédula o nombre y pulse Buscar. Seleccione el cliente de la lista. "
        "Confirme bodega Facturación (es la que tiene las 25 unidades del ejemplo).",
    )
    add_figure(
        doc,
        "12-factura-cliente.png",
        "Figura 11. Cliente elegido. Condición Contado. La bodega Facturación queda fija al agregar líneas.",
    )

    add_step(
        doc,
        10,
        "Agregar la unidad y la caja",
        "En Detalle pulse Buscar artículo. Escriba parte de la descripción (en el ejemplo: «Manual empaque»). "
        "Haga clic en MAN-UND; se agrega 1 unidad. Vuelva a Buscar artículo y elija MAN-CAJ; se agrega 1 caja. "
        "Cada línea lleva su propio precio. El comprobante mostrará ambos códigos.",
    )
    add_figure(
        doc,
        "13-buscador-articulos.png",
        "Figura 12. Buscador: MAN-UND con existencia 25 y MAN-CAJ con 0. El 0 de la caja es el stock propio (vacío); las cajas vendibles son 2.",
    )
    add_figure(
        doc,
        "14-factura-detalle.png",
        "Figura 13. Detalle de la factura: 1 × MAN-UND (₡113) y 1 × MAN-CAJ (₡1.017). Total ₡1.276,90 con impuesto.",
    )

    add_body(
        doc,
        "Al emitir (no se emitió en este ejercicio), SeePOS descarga 1 tableta de MAN-UND por la primera línea "
        "y 10 tabletas de MAN-UND por la caja. Quedarían 14 unidades, equivalentes a 1 caja completa más 4 tabletas sueltas.",
    )

    add_callout(
        doc,
        "Si la clave interna no activa",
        "El banner muestra «Contraseña incorrecta». Confirme que está usando la misma contraseña del ingreso "
        "y que ese cajero tiene una caja abierta. Si la caja no está abierta, Facturación no se habilita aunque la clave sea correcta. "
        "Un administrador puede asignar o restablecer la clave interna del usuario en Parámetros → Usuarios → Credenciales del usuario.",
        fill="FFF4EA",
    )

    # 6
    add_heading_custom(doc, "6. Resumen para el mostrador", 1)
    add_bullet(doc, "Cree primero la unidad, cárguele existencias y luego cree la caja apuntando a esa unidad.", bold_prefix="Alta. ")
    add_bullet(doc, "Otros datos → Artículo a rebajar + Unidades por caja > 1. No dependa del Tipo de presentación.", bold_prefix="Empaque. ")
    add_bullet(doc, "Ajuste, toma y conteo siempre sobre la unidad. La caja solo muestra el cálculo.", bold_prefix="Inventario. ")
    add_bullet(doc, "Active Facturación con la misma contraseña del ingreso. Tiene que haber caja abierta.", bold_prefix="Venta. ")
    add_bullet(doc, "Venda el código que el cliente lleva (unidad o caja). El sistema se encarga de rebajar las unidades.", bold_prefix="Líneas. ")

    add_heading_custom(doc, "7. Qué no hacer", 1)
    add_bullet(doc, "Ajustar existencias o hacer toma física sobre el código de la caja.")
    add_bullet(doc, "Poner como Artículo a rebajar otra caja (no se encadenan empaques).")
    add_bullet(doc, "Dejar Unidades por caja en 1 o en 0: eso no activa el empaque.")
    add_bullet(doc, "Interpretar el 0 de la columna Existencia de la caja en el listado como «no hay cajas para vender».")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(
        "Documento generado a partir de capturas reales de SeePOS. "
        "Los artículos MAN-UND y MAN-CAJ son de demostración para este manual."
    )
    set_run(r, size=9.5, italic=True, color=GRIS)

    doc.save(OUT)
    print("WROTE", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    main()
